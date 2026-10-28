import io
import json
import re

import fitz
from docx import Document
from groq import Groq
from pydantic import BaseModel, Field, ValidationError


MODEL = "openai/gpt-oss-120b"
REQUEST_TIMEOUT_SECONDS = 45.0
MAX_RESUME_CHARS_FOR_ANALYSIS = 10000


class CandidateProfile(BaseModel):
    skills: list[str] = Field(default_factory=list)
    education: list[str] = Field(default_factory=list)
    experience: list[str] = Field(default_factory=list)
    projects: list[str] = Field(default_factory=list)
    certifications: list[str] = Field(default_factory=list)
    target_roles: list[str] = Field(default_factory=list)
    years_of_experience: int = 0


def extract_resume_text(uploaded_file):
    """Extract text from PDF, DOCX, or TXT uploads."""
    filename = uploaded_file.name.lower()
    file_bytes = uploaded_file.getvalue()

    if filename.endswith(".pdf"):
        with fitz.open(stream=file_bytes, filetype="pdf") as document:
            text = "\n".join(page.get_text() for page in document)
    elif filename.endswith(".docx"):
        document = Document(io.BytesIO(file_bytes))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    elif filename.endswith(".txt"):
        text = file_bytes.decode("utf-8", errors="replace")
    else:
        raise ValueError("Please upload a PDF, DOCX, or TXT resume.")

    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        raise ValueError(
            "No readable text was found. If your PDF is scanned, "
            "it may require OCR."
        )
    return text


def parse_resume_with_llm(resume_text):
    """Extract a structured profile from resume text using one API request."""
    if not resume_text or not resume_text.strip():
        raise ValueError("Resume text is empty.")

    normalized_text = re.sub(r"\n{3,}", "\n\n", resume_text).strip()
    if len(normalized_text) > MAX_RESUME_CHARS_FOR_ANALYSIS:
        half = MAX_RESUME_CHARS_FOR_ANALYSIS // 2
        normalized_text = (
            normalized_text[:half]
            + "\n\n[Middle section omitted to reduce token use]\n\n"
            + normalized_text[-half:]
        )

    prompt = f"""
Extract a structured candidate profile from the resume below.

Return one valid JSON object with exactly these fields:
{{
  "skills": ["technical or professional skills"],
  "education": ["education details"],
  "experience": ["employment or internship details"],
  "projects": ["projects explicitly described"],
  "certifications": ["certifications explicitly listed"],
  "target_roles": ["roles supported by the candidate's background"],
  "years_of_experience": 0
}}

Rules:
1. Use only information explicitly supported by the resume.
2. Do not invent skills, qualifications, employers, or projects.
3. Use empty lists when information is unavailable.
4. Use a nonnegative integer for years_of_experience.
5. Include internships in experience only when explicitly listed.
6. Do not count education duration as professional experience.
7. Keep each list item concise while preserving important evidence.
8. Return JSON only, without Markdown fences or commentary.

RESUME:
{normalized_text}
"""

    client = Groq(timeout=REQUEST_TIMEOUT_SECONDS, max_retries=0)
    try:
        response = client.chat.completions.create(
            model=MODEL,
            messages=[
                {
                    "role": "system",
                    "content": (
                        "Extract resume information accurately and concisely. "
                        "Never infer qualifications. Return valid JSON only."
                    ),
                },
                {"role": "user", "content": prompt},
            ],
            response_format={"type": "json_object"},
            temperature=0,
            max_completion_tokens=1600,
        )
        content = response.choices[0].message.content
        if not content:
            raise ValueError("The AI returned an empty resume analysis.")
        data = json.loads(content)
        return CandidateProfile.model_validate(data)
    except (json.JSONDecodeError, ValidationError) as error:
        raise ValueError(
            f"The resume analysis returned invalid profile data: {error}"
        ) from error
    except ValueError:
        raise
    except Exception as error:
        raise RuntimeError(f"Resume analysis request failed: {error}") from error
